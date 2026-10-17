"""المراجعة بقرارين: «سليم» يدخل السند التقرير و«فيه مشكلة» يستبعده؛ الملاحظة والاسم إلزاميان؛ ترتيب السندات بترتيب ظهورها في الملف ثابت."""
import os, sys, io, re, glob
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import openpyxl
import app as A, job_pipeline as JP, report_template as T

if __name__ == "__main__":
    base = os.path.join(T.ROOT, "المنافسات", "منافسة جراحة العظام MED-2026-144")
    bundle = sorted(glob.glob(os.path.join(base, "MED-2026-144-*.pdf")))[3]
    j = JP.create_job("batch", os.path.join(base, "العقد", "عقد المنافسة.pdf"), [bundle], synthetic=True)
    m = JP.run_job(j)
    ws = openpyxl.load_workbook(m["output"])["تقرير السندات"]
    stos = [ws.cell(r, 1).value for r in range(6, ws.max_row + 1) if ws.cell(r, 1).value]
    print("سندات التقرير:", len(stos), "| الحالة:", m["state"])
    # بند مراجعة مصطنع على السند الثالث ثم قرار «فيه مشكلة»
    target = stos[2]
    m = j.load(); m["pending"] = [{"key": "test-key", "reason": "اختبار", "sto": target}]; j.save(m)
    strip = lambda h: re.sub(r"<[^>]+>", "", h)
    print("بلا قرار   :", strip(A.save_decision(j.id, "test-key", None, "", "سبب", "مراجع")))
    print("بلا ملاحظة :", strip(A.save_decision(j.id, "test-key", A.DECISION_BAD, "", "", "مراجع")))
    print("بلا اسم    :", strip(A.save_decision(j.id, "test-key", A.DECISION_BAD, "", "سبب", "")))
    print("صحيح       :", strip(A.save_decision(j.id, "test-key", A.DECISION_BAD, "", "الكمية لا تطابق الفاتورة", "مراجع")))
    m = JP.run_job(j)
    ws = openpyxl.load_workbook(m["output"])["تقرير السندات"]
    stos2 = [ws.cell(r, 1).value for r in range(6, ws.max_row + 1) if ws.cell(r, 1).value]
    print("بعد الاستبعاد:", len(stos2), "| المستبعد غائب:", target not in stos2, "| الترتيب محفوظ:", stos2 == [s for s in stos if s != target])
    print("ملاحظة الرأس:", "استُبعد" in (ws["A2"].value or ""))
    assert target not in stos2 and stos2 == [s for s in stos if s != target] and "استُبعد" in ws["A2"].value
    print("OK")
