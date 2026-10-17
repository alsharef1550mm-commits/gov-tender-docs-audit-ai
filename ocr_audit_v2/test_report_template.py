"""اختبارات قفل قالب Excel: القالب يطابق نفسه، وأي إضافة/إسقاط/إعادة تسمية/ترتيب مختلف يُرفض. تعمل على نسخ مؤقتة فقط."""
import os
import shutil
import tempfile

import openpyxl

import report_template as T


def with_copy(kind, edit):
    td = tempfile.mkdtemp()
    p = os.path.join(td, "t.xlsx")
    shutil.copy(T.template_path(kind), p)
    wb = openpyxl.load_workbook(p)
    edit(wb)
    wb.save(p)
    return p


def test_templates_match_themselves_and_snapshot_is_current():
    for kind in T.TEMPLATE_FILES:
        assert T.template_is_current(kind), kind
        assert T.validate(T.template_path(kind), kind) == [], kind
        T.assert_matches_template(T.template_path(kind), kind, style=False)   # القالب المرجعي القديم بنمطه القديم؛ التقارير الجديدة تُفحص بالنمط


def test_extra_column_header_rejected():
    def edit(wb):
        ws = wb["تقرير السندات"]
        ws.cell(5, ws.max_column + 1, "عمود زائد")
    errs = T.validate(with_copy("batch", edit), "batch")
    assert any("زائدة" in e or "خارج" in e for e in errs), errs


def test_data_column_outside_headers_rejected():
    def edit(wb):
        ws = wb["الأصناف"]
        ws.cell(6, ws.max_column + 2, "بيانات تقنية")
    errs = T.validate(with_copy("batch", edit), "batch")
    assert any("خارج رؤوس" in e for e in errs), errs


def test_dropped_and_renamed_and_reordered_columns_rejected():
    def drop(wb):
        wb["المطابقة"].delete_cols(3)           # يحذف رأس «رقم الفاتورة» وبياناته
    assert T.validate(with_copy("batch", drop), "batch")

    def rename(wb):
        wb["تقرير السندات"].cell(5, 2).value = "تاريخ الاستلام"
    errs = T.validate(with_copy("batch", rename), "batch")
    assert any("زائدة" in e and "ناقصة" in e or "ناقصة" in e for e in errs), errs

    def swap(wb):
        ws = wb["تفصيل الأصناف"]
        a, b = ws.cell(5, 3).value, ws.cell(5, 4).value
        ws.cell(5, 3).value, ws.cell(5, 4).value = b, a
    errs = T.validate(with_copy("competition", swap), "competition")
    assert any("ترتيب" in e for e in errs), errs


def test_sheet_changes_rejected():
    def add_sheet(wb):
        wb.create_sheet("بيانات تقنية")
    assert any("زائدة" in e for e in T.validate(with_copy("competition", add_sheet), "competition"))

    def rename_sheet(wb):
        wb["المطابقة"].title = "مطابقة"
    assert T.validate(with_copy("batch", rename_sheet), "batch")

    def reorder(wb):
        wb.move_sheet("تقرير المنافسة والأصناف", offset=-3)
    errs = T.validate(with_copy("competition", reorder), "competition")
    assert any("ترتيب الشيتات" in e for e in errs), errs


def test_assert_raises_on_violation():
    p = with_copy("batch", lambda wb: wb.create_sheet("x"))
    try:
        T.assert_matches_template(p, "batch")
    except ValueError as e:
        assert "يخالف القالب" in str(e)
    else:
        raise AssertionError("كان يجب أن يرفض")


if __name__ == "__main__":
    for n, f in list(globals().items()):
        if n.startswith("test_"):
            f()
            print("ok", n)
