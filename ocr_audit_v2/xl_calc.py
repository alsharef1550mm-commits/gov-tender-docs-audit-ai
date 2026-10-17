"""حاسب صيغ Excel الصغير لتقارير الدفعات والمنافسات: يحسب قيم الصيغ ويحفظها كنتائج مخزّنة (cached) داخل ملف xlsx
حتى تظهر الأرقام فور الفتح في أي عارض (Excel يعيد الحساب عند التحميل أيضاً).

يغطي الدوال المستعملة فعلاً في القالبين: ROUND, SUM, SUMIF, SUMIFS, COUNTIF, IF, MAX, AND (وMIN, OR, ABS احتياطاً)،
والعمليات + - * / ^ & ومقارنات = <> < > <= >= والنسبة %. أي دالة غير معروفة تُرفع صراحة ولا تُخمَّن.
يُتحقق منه مقابل نتائج مخزّنة في ملفات حسبها Excel/LibreOffice (انظر test_xl_calc.py).

الاستعمال: add_cached_values(path)  # يعدّل الملف في مكانه
"""
from __future__ import annotations

import datetime as dt
import fnmatch
import os
import re
import shutil
import tempfile
import zipfile
from decimal import ROUND_HALF_UP, Decimal
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
EPOCH = dt.datetime(1899, 12, 30)

TOKEN = re.compile(r"""
    (?P<ws>\s+)
  | (?P<str>"(?:[^"]|"")*")
  | (?P<ref>(?:(?:'(?:[^']|'')+'|[A-Za-z_؀-ۿ][\w؀-ۿ\.]*)!)?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)
  | (?P<num>\d+(?:\.\d+)?(?:[eE][+-]?\d+)?)
  | (?P<func>[A-Z][A-Z0-9\.]*)(?=\()
  | (?P<bool>TRUE|FALSE)
  | (?P<op><>|<=|>=|[-+*/^&=<>%(),])
""", re.X)


class CalcError(Exception):
    def __init__(self, code):
        super().__init__(code)
        self.code = code


def _num(x):
    if x is None or x == "":
        return 0.0
    if isinstance(x, bool):
        return 1.0 if x else 0.0
    if isinstance(x, (int, float)):
        return float(x)
    if isinstance(x, Decimal):
        return float(x)
    if isinstance(x, dt.datetime):
        return (x - EPOCH).total_seconds() / 86400
    if isinstance(x, dt.date):
        return float((dt.datetime(x.year, x.month, x.day) - EPOCH).days)
    try:
        return float(str(x).replace(",", ""))
    except ValueError:
        raise CalcError("#VALUE!")


def _cmp_norm(a, b):
    """يوحّد طرفي المقارنة: الفارغ يساوي 0 مع الأرقام و"" مع النص؛ النص لا يطابق الرقم."""
    if a is None:
        a = "" if isinstance(b, str) else (False if isinstance(b, bool) else 0.0)
    if b is None:
        b = "" if isinstance(a, str) else (False if isinstance(a, bool) else 0.0)
    if isinstance(a, (dt.date, dt.datetime, Decimal)):
        a = _num(a)
    if isinstance(b, (dt.date, dt.datetime, Decimal)):
        b = _num(b)
    return a, b


def _compare(op, a, b):
    a, b = _cmp_norm(a, b)
    if isinstance(a, str) and isinstance(b, str):
        a, b = a.lower(), b.lower()
    elif isinstance(a, str) != isinstance(b, str):
        # نص مقابل رقم: في Excel النص دائماً أكبر من الرقم؛ المساواة خاطئة
        if op == "=":
            return False
        if op == "<>":
            return True
        sa = 1 if isinstance(a, str) else 0
        sb = 1 if isinstance(b, str) else 0
        a, b = sa, sb
    return {"=": a == b, "<>": a != b, "<": a < b, ">": a > b, "<=": a <= b, ">=": a >= b}[op]


def _round(x, n):
    n = int(n)
    q = Decimal(1).scaleb(-n)
    return float(Decimal(repr(float(x))).quantize(q, rounding=ROUND_HALF_UP))


def _crit(c):
    """معيار COUNTIF/SUMIF -> دالة تطابق قيمة الخلية."""
    if isinstance(c, str):
        m = re.match(r"^(<>|<=|>=|=|<|>)(.*)$", c)
        op, rest = (m.group(1), m.group(2)) if m else ("=", c)
        try:
            val = float(rest)
        except ValueError:
            val = rest
        if isinstance(val, str) and op in ("=", "<>") and any(ch in val for ch in "*?"):
            pat = val.lower()
            return lambda v: (isinstance(v, str) and fnmatch.fnmatchcase(v.lower(), pat)) == (op == "=")
        if val == "" and op == "=":
            return lambda v: v is None or v == ""

        def f(v):
            if v is None:
                return op == "<>" if val != "" else op == "="
            if isinstance(val, float):
                if isinstance(v, (int, float, Decimal)) and not isinstance(v, bool):
                    return _compare(op, float(v), val)
                return op == "<>"
            return _compare(op, v if isinstance(v, str) else str(v), val) if isinstance(v, str) else (op == "<>")
        return f
    cv = c
    if cv is None:
        return lambda v: v is None or v == 0
    return lambda v: v is not None and not isinstance(v, str) and float(_num(v)) == float(_num(cv)) if not isinstance(cv, str) else False


class Book:
    def __init__(self, wb):
        self.wb = wb
        self.cache = {}
        self.stack = set()
        self.errors = []

    # ------------------------------------------------------------ قراءة الخلايا
    def cell(self, sheet, row, col):
        key = (sheet, row, col)
        if key in self.cache:
            return self.cache[key]
        ws = self.wb[sheet]
        v = ws.cell(row, col).value
        if isinstance(v, str) and v.startswith("="):
            if key in self.stack:
                raise CalcError("#REF!")
            self.stack.add(key)
            try:
                v = Parser(self, sheet, v[1:]).run()
            except CalcError as e:
                v = ("#ERR", e.code)
            finally:
                self.stack.discard(key)
        self.cache[key] = v
        return v

    def rng(self, sheet, r1, c1, r2, c2):
        ws = self.wb[sheet]
        r2 = min(r2, max(ws.max_row, r1))
        c2 = min(c2, max(ws.max_column, c1))
        return [[self.cell(sheet, r, c) for c in range(c1, c2 + 1)] for r in range(r1, r2 + 1)]


class Rng(list):
    """نطاق خلايا (قائمة صفوف)."""


class Parser:
    def __init__(self, book, sheet, text):
        self.book, self.sheet = book, sheet
        self.toks = []
        pos = 0
        while pos < len(text):
            m = TOKEN.match(text, pos)
            if not m:
                raise ValueError(f"رمز غير مفهوم في الصيغة عند: {text[pos:pos + 20]!r}")
            pos = m.end()
            kind = m.lastgroup
            if kind != "ws":
                self.toks.append((kind, m.group(kind)))
        self.i = 0

    # ---- أدوات
    def peek(self):
        return self.toks[self.i] if self.i < len(self.toks) else (None, None)

    def take(self):
        t = self.peek()
        self.i += 1
        return t

    def expect(self, op):
        t = self.take()
        if t != ("op", op):
            raise ValueError(f"المتوقع {op} وجدنا {t}")

    def run(self):
        v = self.expr()
        if self.i != len(self.toks):
            raise ValueError("بقايا غير مفسَّرة في الصيغة")
        return self.scalar(v)

    @staticmethod
    def scalar(v):
        if isinstance(v, Rng):
            return v[0][0] if v and v[0] else None
        return v

    # ---- قواعد النحو (الأولوية: مقارنة < & < جمع < ضرب < أس < أحادي < نسبة)
    def expr(self):
        a = self.concat()
        while self.peek()[0] == "op" and self.peek()[1] in ("=", "<>", "<", ">", "<=", ">="):
            op = self.take()[1]
            b = self.concat()
            a = _compare(op, self.scalar(a), self.scalar(b))
        return a

    def concat(self):
        a = self.add()
        while self.peek() == ("op", "&"):
            self.take()
            b = self.add()
            a = self.txt(self.scalar(a)) + self.txt(self.scalar(b))
        return a

    @staticmethod
    def txt(v):
        if v is None:
            return ""
        if isinstance(v, float) and v == int(v):
            return str(int(v))
        return str(v)

    def add(self):
        a = self.mul()
        while self.peek()[0] == "op" and self.peek()[1] in "+-" and self.peek()[1] != "":
            op = self.take()[1]
            b = self.mul()
            x, y = self.n(a), self.n(b)
            a = x + y if op == "+" else x - y
        return a

    def mul(self):
        a = self.pow()
        while self.peek()[0] == "op" and self.peek()[1] in ("*", "/"):
            op = self.take()[1]
            b = self.pow()
            x, y = self.n(a), self.n(b)
            if op == "/":
                if y == 0:
                    raise CalcError("#DIV/0!")
                a = x / y
            else:
                a = x * y
        return a

    def pow(self):
        a = self.unary()
        while self.peek() == ("op", "^"):
            self.take()
            b = self.unary()
            a = self.n(a) ** self.n(b)
        return a

    def unary(self):
        if self.peek() == ("op", "-"):
            self.take()
            return -self.n(self.unary())
        if self.peek() == ("op", "+"):
            self.take()
            return self.unary()
        return self.postfix()

    def postfix(self):
        a = self.atom()
        while self.peek() == ("op", "%"):
            self.take()
            a = self.n(a) / 100
        return a

    def n(self, v):
        v = self.scalar(v)
        if isinstance(v, tuple) and v and v[0] == "#ERR":
            raise CalcError(v[1])
        return _num(v)

    def atom(self):
        kind, val = self.take()
        if kind == "num":
            return float(val)
        if kind == "str":
            return val[1:-1].replace('""', '"')
        if kind == "bool":
            return val == "TRUE"
        if kind == "ref":
            return self.ref(val)
        if kind == "func":
            return self.call(val)
        if (kind, val) == ("op", "("):
            v = self.expr()
            self.expect(")")
            return v
        raise ValueError(f"رمز غير متوقع {kind} {val}")

    def ref(self, text):
        sheet = self.sheet
        if "!" in text:
            sh, text = text.rsplit("!", 1)
            sheet = sh[1:-1].replace("''", "'") if sh.startswith("'") else sh
        text = text.replace("$", "")
        parts = text.split(":")

        def rc(p):
            m = re.match(r"([A-Z]{1,3})(\d+)$", p)
            return int(m.group(2)), column_index_from_string(m.group(1))
        r1, c1 = rc(parts[0])
        if len(parts) == 1:
            v = self.book.cell(sheet, r1, c1)
            if isinstance(v, tuple) and v and v[0] == "#ERR":
                raise CalcError(v[1])
            return v
        r2, c2 = rc(parts[1])
        return Rng(self.book.rng(sheet, min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2)))

    # ---- الدوال
    def args_lazy(self):
        """يقرأ وسائط الدالة كمواضع رموز ليُقيَّم كل منها عند الحاجة (IF كسولة)."""
        self.expect("(")
        starts = []
        depth = 0
        cur = self.i
        if self.peek() == ("op", ")"):
            self.take()
            return []
        while True:
            t = self.peek()
            if t[0] is None:
                raise ValueError("قوس غير مغلق")
            if t == ("op", "("):
                depth += 1
            elif t == ("op", ")"):
                if depth == 0:
                    starts.append((cur, self.i))
                    self.take()
                    return starts
                depth -= 1
            elif t == ("op", ",") and depth == 0:
                starts.append((cur, self.i))
                cur = self.i + 1
            self.take()

    def eval_range(self, span):
        sub = Parser.__new__(Parser)
        sub.book, sub.sheet, sub.toks, sub.i = self.book, self.sheet, self.toks[span[0]:span[1]], 0
        v = sub.expr()
        if sub.i != len(sub.toks):
            raise ValueError("وسيط غير مفسَّر")
        return v

    @staticmethod
    def flat(vals, direct_ok=True):
        out = []
        for v in vals:
            if isinstance(v, Rng):
                for row in v:
                    out.extend(x for x in row if isinstance(x, (int, float, Decimal)) and not isinstance(x, bool))
            elif v is not None and v != "":
                out.append(_num(v))
        return out

    @staticmethod
    def cells(rng):
        return [x for row in rng for x in row]

    def call(self, name):
        spans = self.args_lazy()
        ev = lambda k: self.eval_range(spans[k])
        if name == "IF":
            c = self.scalar(ev(0))
            if isinstance(c, tuple) and c and c[0] == "#ERR":
                raise CalcError(c[1])
            truth = bool(c) if not isinstance(c, str) else (_ for _ in ()).throw(CalcError("#VALUE!"))
            if truth:
                return ev(1) if len(spans) > 1 else True
            return ev(2) if len(spans) > 2 else False
        vals = [ev(k) for k in range(len(spans))]
        for v in vals:
            sv = self.scalar(v) if not isinstance(v, Rng) else None
            if isinstance(sv, tuple) and sv and sv[0] == "#ERR":
                raise CalcError(sv[1])
        if name == "ROUND":
            return _round(self.n(vals[0]), self.n(vals[1]))
        if name == "SUM":
            return float(sum(self.flat(vals)))
        if name == "MAX":
            f = self.flat(vals)
            return max(f) if f else 0.0
        if name == "MIN":
            f = self.flat(vals)
            return min(f) if f else 0.0
        if name == "ABS":
            return abs(self.n(vals[0]))
        if name in ("AND", "OR"):
            bs = [bool(self.scalar(v)) if not isinstance(v, Rng) else any(bool(x) for x in self.cells(v)) for v in vals]
            return all(bs) if name == "AND" else any(bs)
        if name == "COUNTIF":
            m = _crit(self.scalar(vals[1]))
            return float(sum(1 for x in self.cells(vals[0]) if m(x)))
        if name == "SUMIF":
            m = _crit(self.scalar(vals[1]))
            crit_cells = self.cells(vals[0])
            sum_cells = self.cells(vals[2]) if len(vals) > 2 else crit_cells
            tot = 0.0
            for x, s in zip(crit_cells, sum_cells):
                if m(x) and isinstance(s, (int, float, Decimal)) and not isinstance(s, bool):
                    tot += float(s)
            return tot
        if name == "SUMIFS":
            sum_cells = self.cells(vals[0])
            ms = [(self.cells(vals[k]), _crit(self.scalar(vals[k + 1]))) for k in range(1, len(vals), 2)]
            tot = 0.0
            for idx, s in enumerate(sum_cells):
                if all(m(cs[idx]) for cs, m in ms if idx < len(cs)) and isinstance(s, (int, float, Decimal)) and not isinstance(s, bool):
                    tot += float(s)
            return tot
        raise ValueError(f"دالة غير مدعومة في حاسب التقارير: {name}")


# ---------------------------------------------------------------------- واجهة
def compute(path):
    """{اسم الشيت: {إحداثي الخلية: قيمة}} لكل خلايا الصيغ في الملف."""
    wb = openpyxl.load_workbook(path)
    book = Book(wb)
    out = {}
    for ws in wb:
        d = {}
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    d[c.coordinate] = book.cell(ws.title, c.row, c.column)
        out[ws.title] = d
    return out


def _sheet_files(zf):
    wbx = ET.fromstring(zf.read("xl/workbook.xml"))
    rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    rid2t = {r.get("Id"): r.get("Target") for r in rels}
    res = {}
    for s in wbx.find(f"{{{NS_MAIN}}}sheets"):
        t = rid2t[s.get(f"{{{NS_REL}}}id")]
        res[s.get("name")] = ("xl/" + t.lstrip("/")) if not t.startswith("/") else t.lstrip("/")
        res[s.get("name")] = res[s.get("name")].replace("xl/xl/", "xl/")
    return res


def add_cached_values(path):
    """يكتب نتيجة كل صيغة داخل <v> في الملف نفسه. يعيد عدد الصيغ المحفوظة وقائمة الأخطاء (تظهر كأخطاء Excel)."""
    values = compute(path)
    errors = []
    ET.register_namespace("", NS_MAIN)
    ET.register_namespace("r", NS_REL)
    ET.register_namespace("mc", "http://schemas.openxmlformats.org/markup-compatibility/2006")
    ET.register_namespace("x14ac", "http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac")
    tmp = tempfile.mktemp(suffix=".xlsx")
    n = 0
    with zipfile.ZipFile(path) as zin:
        files = _sheet_files(zin)
        by_file = {f: s for s, f in files.items()}
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename in by_file:
                    sheet = by_file[item.filename]
                    root = ET.fromstring(data)
                    for c in root.iter(f"{{{NS_MAIN}}}c"):
                        f = c.find(f"{{{NS_MAIN}}}f")
                        if f is None:
                            continue
                        val = values[sheet].get(c.get("r"))
                        v = c.find(f"{{{NS_MAIN}}}v")
                        if v is None:
                            v = ET.SubElement(c, f"{{{NS_MAIN}}}v")
                        if isinstance(val, tuple) and val and val[0] == "#ERR":
                            c.set("t", "e")
                            v.text = val[1]
                            errors.append((sheet, c.get("r"), val[1]))
                        elif isinstance(val, bool):
                            c.set("t", "b")
                            v.text = "1" if val else "0"
                        elif isinstance(val, str):
                            c.set("t", "str")
                            v.text = val
                        elif val is None:
                            c.set("t", "str")
                            v.text = ""
                        else:
                            if "t" in c.attrib:
                                del c.attrib["t"]
                            v.text = repr(float(val)) if float(val) != int(float(val)) or abs(val) > 1e15 else str(int(float(val)))
                        n += 1
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                zout.writestr(item, data)
    shutil.move(tmp, path)
    return n, errors
