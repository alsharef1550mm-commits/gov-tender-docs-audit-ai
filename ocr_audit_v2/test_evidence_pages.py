"""اختبار أدلة الصفحات في تقرير الدفعة (305-01-001): لكل صف صنف، صفحة السند والفاتورة والتعميد يجب أن تحمل سطر الصنف نفسه
(الرمز، وللسند والفاتورة الكمية، وللثلاثة السعر)، وسطر الفاتورة يقع تحت رقم السند الصحيح (آخر STO مطبوع قبله).

REPORT_IN=<تقرير سابق> يفحص تقريراً قائماً بدل توليد جديد (للمقارنة بين نسختين).
الخطآن اللذان رصدهما المالك: صفحة فاتورة LAB-008 للسند STO0000002440 (62 بدل 63) وصفحة التعميد 7 (بداية التعميد) بدل 8 (سطر الصنف).
"""
import glob
import io
import os
import re
import sys
import tempfile

import openpyxl
import pdfplumber

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import report_template as T  # noqa: E402
import report_writer as RW  # noqa: E402

RE_STO = re.compile(r"STO\d{10}")


def _text(pdf, page):
    try:
        return pdf.pages[int(page) - 1].extract_text() or ""
    except Exception:  # noqa: BLE001
        return ""


def _nums(line):
    return [float(t) for t in line.replace(",", "").split() if t.replace(".", "", 1).isdigit()]


def line_ok(pdf, page, code, qty, price, sto=None):
    """سطر الصنف في الصفحة يحمل الكمية والسعر (إن أُعطيا)، وإن أُعطي sto فيقع تحت ذلك السند."""
    cur = None
    for ln in _text(pdf, page).splitlines():
        m = RE_STO.search(ln)
        if m:
            cur = m.group(0)
        if code not in ln:
            continue
        nums = _nums(ln)
        if qty is not None and not any(abs(n - float(qty)) < 1e-9 for n in nums):
            continue
        if price is not None and not any(abs(n - float(price)) < 0.006 for n in nums):
            continue
        if sto is not None and cur != sto:
            continue
        return True
    return False


def main():
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    folder = glob.glob(os.path.join(T.ROOT, "المنافسات", "*MED-2026-305"))[0]
    bundle = os.path.join(folder, "MED-2026-305-01-001.pdf")
    contract = glob.glob(os.path.join(folder, "العقد", "*.pdf"))[0]
    out = os.environ.get("REPORT_IN")
    if not out:
        out = os.path.join(tempfile.mkdtemp(), "تقرير فحص الدفعة MED-2026-305-01-001.xlsx")
        RW.write_batch_report(bundle, contract, out)
    ws = openpyxl.load_workbook(out, data_only=True)["الأصناف"]
    col = {c.value: i + 1 for i, c in enumerate(ws[5]) if c.value}
    pdf = pdfplumber.open(bundle)
    bad = []
    rows = 0
    for r in range(6, ws.max_row + 1):
        code = ws.cell(r, col["رقم الصنف"]).value
        if not code:
            continue
        rows += 1
        sto, qty, price = ws.cell(r, 1).value, ws.cell(r, col["عدد الوحدات"]).value, ws.cell(r, col["سعر الوحدة"]).value
        checks = (("صفحة السند", qty, None), ("صفحة الفاتورة", qty, sto), ("صفحة التعميد", None, None))
        for name, q, own in checks:
            pg = ws.cell(r, col[name]).value
            if pg and not line_ok(pdf, pg, code, q, price, own):
                bad.append((name, sto, code, pg))
    print(f"صفوف الأصناف {rows} | مراجع صفحات لا تحمل سطر الصنف: {len(bad)}")
    for b in bad[:10]:
        print("  ", b)
    print("OK" if not bad else f"FAIL {len(bad)}")
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
