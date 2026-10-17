"""تنسيق التقارير الجديدة (توجيه صاحب المشروع، 16 أكتوبر 2026): نمط المراجعة المعتمد من النسخة المرجعية
«تقرير فحص الدفعة MED-2026-305-01-001 - منسق للمراجعة.xlsx».

ما يتغير (تنسيق فقط؛ لا اسم شيت ولا عنوان عمود ولا ترتيب ولا قيمة ولا صيغة):
  1) رؤوس الجداول: خلفية أزرق فاتح E1EFFA وخط داكن غامق 17324D وحدود رفيعة؛ ارتفاع الصف 82.
  2) حدود رفيعة (لون 9BB3C9) لكل خلايا الرأس والبيانات حتى آخر صف بيانات، بما فيها الخلايا الفارغة المقصودة.
  3) سطور العنوان والمصدر والملاحظات أعلى كل شيت (التي لا تحمل غير العمود A) تُدمج A:H: العنوان EFF5FB بخط 15 غامق،
     والمصدر/الملاحظات F7FAFD بخط 10 رمادي، ارتفاعاهما 38 و34.
  4) نتائج الصيغ مخزّنة (xl_calc.add_cached_values) فتظهر الأرقام فور الفتح.
لا يلمس القوالب المصدرية ولا التقارير القديمة؛ يُطبَّق على المصنف الجديد قبل حفظه فقط (finalize).
"""
from __future__ import annotations

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import xl_calc

HEADER_FILL = "FFE1EFFA"
HEADER_FONT = "FF17324D"
TITLE_FILL = "FFEFF5FB"
TITLE_FONT = "FF17365D"
NOTE_FILL = "FFF7FAFD"
NOTE_FONT = "FF5B6778"
BORDER_RGB = "FF9BB3C9"
MERGE_COLS = 8                 # A:H
H_TITLE, H_NOTE, H_HEADER = 38, 34, 82

SIDE = Side(style="thin", color=BORDER_RGB)
BORDER = Border(left=SIDE, right=SIDE, top=SIDE, bottom=SIDE)


def _solid(rgb):
    return PatternFill("solid", fgColor=rgb, bgColor=rgb)


def header_rows(ws):
    """صفوف رؤوس الجداول مستقلة عن لون التعبئة: خلية A غامقة بتعبئة صلبة بخط ≤12 وغير مدمجة."""
    wide = {r for rng in ws.merged_cells.ranges if rng.min_col == 1 and rng.max_col > 1 for r in range(rng.min_row, rng.max_row + 1)}
    rows = []
    for r in range(1, ws.max_row + 1):
        c = ws.cell(r, 1)
        if r in wide:
            continue
        if c.font is not None and c.font.b and c.fill is not None and c.fill.fill_type == "solid" and (c.font.sz or 10) <= 12 and c.value is not None:
            vals = [ws.cell(r, k).value for k in range(1, ws.max_column + 1)]
            if all(v is None or (isinstance(v, str) and not v.startswith("=")) for v in vals):      # صف الإجمالي يحمل صيغاً/أرقاماً فليس رأساً
                rows.append(r)
    return rows


def header_width(ws, r):
    w = ws.max_column
    while w > 0 and ws.cell(r, w).value is None:
        w -= 1
    return w


def _top_rows(ws, first_header):
    """سطور أعلى الشيت التي لا تحمل غير العمود A (عنوان/مصدر/ملاحظة)."""
    out = []
    for r in range(1, first_header):
        a = ws.cell(r, 1).value
        rest = [ws.cell(r, c).value for c in range(2, ws.max_column + 1)]
        if a is not None and all(v is None for v in rest):
            out.append(r)
    return out


def apply_style(wb):
    """يطبّق نمط المراجعة على المصنف في الذاكرة (يعيد ملخصاً بما طُبّق)."""
    summary = {}
    for ws in wb:
        hdr = header_rows(ws)
        if not hdr:
            summary[ws.title] = "بلا رؤوس"
            continue
        ncols_merge = min(MERGE_COLS, max(ws.max_column, 1))
        for k, r in enumerate(_top_rows(ws, hdr[0])):
            c = ws.cell(r, 1)
            is_title = k == 0
            c.fill = _solid(TITLE_FILL if is_title else NOTE_FILL)
            c.font = Font(name="Arial", b=is_title, sz=15 if is_title else 10, color=TITLE_FONT if is_title else NOTE_FONT)
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            for cc in range(1, ncols_merge + 1):
                ws.cell(r, cc).border = BORDER
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols_merge)
            ws.row_dimensions[r].height = H_TITLE if is_title else H_NOTE
        for i, h in enumerate(hdr):
            w = header_width(ws, h)
            for c in range(1, w + 1):
                cell = ws.cell(h, c)
                cell.fill = _solid(HEADER_FILL)
                cell.font = Font(name="Arial", b=True, sz=10, color=HEADER_FONT)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = BORDER
            ws.row_dimensions[h].height = max(H_HEADER, ws.row_dimensions[h].height or 0)
            end = (hdr[i + 1] - 1) if i + 1 < len(hdr) else ws.max_row
            last = h
            for r in range(h + 1, end + 1):
                if any(ws.cell(r, c).value is not None for c in range(1, w + 1)):
                    last = r
            for r in range(h + 1, last + 1):
                for c in range(1, w + 1):
                    cell = ws.cell(r, c)
                    cell.border = BORDER
                    al = cell.alignment
                    cell.alignment = Alignment(horizontal=al.horizontal or "center", vertical="center", wrap_text=al.wrap_text, indent=al.indent)
        summary[ws.title] = f"{len(hdr)} كتلة رؤوس"
    for ws in wb:                                              # لا نص أبيض في أي تقرير جديد
        for row in ws.iter_rows():
            for c in row:
                if _is_white(c):
                    f = c.font
                    c.font = Font(name=f.name, sz=f.sz, b=f.b, i=f.i, color=HEADER_FONT)
                    if c.fill.fill_type == "solid":
                        c.fill = _solid(HEADER_FILL)
    return summary


def _is_white(c):
    col = c.font.color
    return col is not None and col.type == "rgb" and isinstance(col.rgb, str) and col.rgb.upper() in ("FFFFFFFF", "00FFFFFF")


def finalize(wb, out_path):
    """تنسيق المراجعة ثم الحفظ ثم تخزين نتائج الصيغ. يعيد (عدد الصيغ المحفوظة، أخطاء الصيغ)."""
    import os
    apply_style(wb)
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    n, errors = xl_calc.add_cached_values(out_path)
    return n, errors


def validate_style(path):
    """مخالفات نمط المراجعة في تقرير جديد؛ فارغة = مطابق. لا تُستعمل على التقارير القديمة."""
    import zipfile
    from xml.etree import ElementTree as ET
    errs = []
    wb = openpyxl.load_workbook(path)
    for ws in wb:
        hdr = header_rows(ws)
        if not hdr:
            continue
        for r in _top_rows(ws, hdr[0]):
            if not any(r == m.min_row and m.min_col == 1 and m.max_col >= min(MERGE_COLS, ws.max_column) for m in ws.merged_cells.ranges):
                errs.append(f"[{ws.title}] سطر {r} (عنوان/مصدر/ملاحظة) غير مدمج A:H")
        for i, h in enumerate(hdr):
            w = header_width(ws, h)
            for c in range(1, w + 1):
                cell = ws.cell(h, c)
                f = cell.fill
                if not (f.fill_type == "solid" and f.fgColor.rgb == HEADER_FILL):
                    errs.append(f"[{ws.title}] رأس {cell.coordinate} بلا خلفية الأزرق الفاتح المعتمدة")
                    break
                if not (cell.font.b and cell.font.color is not None and cell.font.color.rgb == HEADER_FONT):
                    errs.append(f"[{ws.title}] رأس {cell.coordinate} بلا الخط الداكن الغامق المعتمد")
                    break
            end = (hdr[i + 1] - 1) if i + 1 < len(hdr) else ws.max_row
            last = h
            for r in range(h + 1, end + 1):
                if any(ws.cell(r, c).value is not None for c in range(1, w + 1)):
                    last = r
            bad = [ws.cell(r, c).coordinate for r in range(h, last + 1) for c in range(1, w + 1)
                   if not all(getattr(ws.cell(r, c).border, s).style == "thin" for s in ("left", "right", "top", "bottom"))]
            if bad:
                errs.append(f"[{ws.title}] {len(bad)} خلية رأس/بيانات بلا حدود رفيعة كاملة (أولها {bad[0]})")
    for ws in wb:
        white = [c.coordinate for row in ws.iter_rows() for c in row if _is_white(c)]
        if white:
            errs.append(f"[{ws.title}] {len(white)} خلية بنص أبيض (أولها {white[0]})؛ ممنوع في التقارير الجديدة")
    # نتائج الصيغ المخزّنة: كل <f> يلزمه <v>
    ns = "{" + xl_calc.NS_MAIN + "}"
    with zipfile.ZipFile(path) as zf:
        files = xl_calc._sheet_files(zf)
        for sheet, fn in files.items():
            root = ET.fromstring(zf.read(fn))
            miss = 0
            for c in root.iter(ns + "c"):
                if c.find(ns + "f") is not None:
                    v = c.find(ns + "v")
                    if v is None or (v.text in (None, "") and c.get("t") != "str"):
                        miss += 1
            if miss:
                errs.append(f"[{sheet}] {miss} صيغة بلا نتيجة محفوظة")
    return errs
