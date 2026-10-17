"""كاتب تقرير المنافسة وفق قالب «تقرير فحص المنافسة MED-2026-339.xlsx» حرفياً (الأسماء والأعمدة والصيغ).

- يجمّع كل ربطات المنافسة، ويطابق الفاتورة كـ«مبرر مؤيد» على مستوى **المنافسة كلها**: كل السندات التي تذكر رقم الفاتورة عبر
  الربطات والتعاميد تُجمع قبل نسب البنود، ثم تُوزَّع المبالغ على سنداتها بلا تكرار (invoice_match.py).
- ما لا يُحسم يبقى خلية فارغة ويُسجَّل في قائمة unresolved مع صفحة الدليل؛ لا يُعلن إقفال: خلية الإقفال تصبح «يتطلب مراجعة».
- أعمدة العقد (المبلغ في العقد، إجمالي الدفعة في العقد) تُقرأ من نص العقد (contract_data.py) أو تُترك فارغة.
"""
from __future__ import annotations

import copy
import datetime as dt
import os
import re
from collections import defaultdict
from decimal import Decimal as D
from multiprocessing import Pool

import openpyxl

import contract_data as CD
import engine
import extract
import invoice_match as IM
import report_style as RS
import report_template as T
from report_writer import _pick_after, _runs, _same_row, _save, review_item

SEP = "، "


def _date(s):
    try:
        return dt.datetime.strptime(s, "%Y-%m-%d")
    except Exception:  # noqa: BLE001
        return None


def _f(x):
    return None if x is None else float(x)


def _fmt_qty(q):
    return int(q) if q == q.to_integral_value() else float(q)


def _work(args):
    path, rules, cache_dir = args
    try:
        cp = None
        if cache_dir:
            import hashlib
            h = hashlib.sha256()
            with open(path, "rb") as f:
                for chunk in iter(lambda: f.read(1 << 20), b""):
                    h.update(chunk)
            cp = os.path.join(cache_dir, h.hexdigest()[:24] + ".pkl")
            if os.path.exists(cp):
                import pickle
                with open(cp, "rb") as f:
                    return pickle.load(f)                 # استئناف: نتيجة الربطة محفوظة
        b = extract.read_bundle(path)
        out = {"path": path, "bundle": b, "results": engine.check_bundle(b, rules)}
        if cp:
            import pickle
            with open(cp + ".tmp", "wb") as f:
                pickle.dump(out, f)
            os.replace(cp + ".tmp", cp)
        return out
    except Exception as e:  # noqa: BLE001
        return {"path": path, "error": f"{type(e).__name__}: {e}"}


def collect_competition(bundle_paths, contract_path, code_map=None, workers=1, progress=None, cache_dir=None):
    code_map = code_map or {}
    first = os.path.basename(sorted(bundle_paths)[0])
    mm = re.match(r"(.+)-\d{2}-\d{3}\.pdf$", first)
    comp_no = mm.group(1) if mm else os.path.splitext(first)[0]      # الاسم لا يشترط صيغة معينة: عند عدم المطابقة يُستعمل اسم الملف
    rules = engine.read_rules(contract_path, comp_no)
    contract = CD.read_contract(contract_path)
    paths = sorted(bundle_paths)
    if workers > 1:
        with Pool(workers) as pool:
            outs = []
            for i, o in enumerate(pool.imap(_work, [(p, rules, cache_dir) for p in paths]), 1):
                outs.append(o)
                if progress:
                    progress(i, len(paths), os.path.basename(o["path"]))
    else:
        outs = []
        for i, p in enumerate(paths, 1):
            outs.append(_work((p, rules, cache_dir)))
            if progress:
                progress(i, len(paths), os.path.basename(p))
    unresolved = []
    errors = [o for o in outs if "error" in o]
    for e in errors:
        unresolved.append({"reason": "تعذّرت قراءة الربطة", "file": os.path.basename(e["path"]), "detail": e["error"]})
    outs = [o for o in outs if "error" not in o]

    receipts, results, bundle_of, order = {}, {}, {}, []
    dups = []
    pos, po_first_bundle = {}, {}
    inv_by_bundle = {}
    for o in outs:
        b = o["bundle"]
        res_by = {}
        for r in o["results"]:
            res_by.setdefault(r.sto, r)                     # النسخة المكررة داخل الربطة لا تحل محل الأولى
        inv_by_bundle[b.file] = {v.number: v for v in b.invoices}
        for r in b.receipts:
            if r.sto in receipts:
                dups.append((r.sto, b.file, r))             # يُحتسب مرة واحدة (لا يتضاعف المبلغ) لكن يُبلَّغ عنه: احتمال ازدواج السداد
                continue
            receipts[r.sto] = r
            results[r.sto] = res_by[r.sto]
            bundle_of[r.sto] = b.file
            order.append(r.sto)
        for p in b.pos:
            if p.number not in pos:
                pos[p.number] = p
                po_first_bundle[p.number] = b.file
    all_res = [results[s] for s in order]
    engine.apply_cumulative_cap(all_res, rules)

    # ---- مطابقة الفاتورة على مستوى المنافسة
    copies_of = defaultdict(list)                           # رقم الفاتورة -> نسخ (من كل الربطات)
    for o in outs:
        b = o["bundle"]
        for v in b.invoices:
            by_copy = defaultdict(list)
            for it in v.items:
                by_copy[it.copy].append({"code": it.code, "qty": it.qty, "price": it.unit_price, "amount": it.total, "net": it.net,
                                         "nupco": it.nupco, "vat": it.vat, "page": it.ev.page, "file": b.file, "expiry": it.expiry})
            for c in sorted(by_copy):
                copies_of[v.number].append(by_copy[c])
    cit = defaultdict(list)
    for s in order:
        if receipts[s].invoice_no:
            cit[receipts[s].invoice_no].append(receipts[s])
    attribution = {}
    for num, rs in cit.items():
        if num not in copies_of:
            for r in rs:
                unresolved.append({"reason": "السند يذكر فاتورة غير موجودة في ربطات المنافسة", "sto": r.sto, "invoice": num,
                                   "page": r.pages[0] if r.pages else None, "file": bundle_of[r.sto]})
            continue
        out = IM.match_invoice(
            [{"sto": r.sto, "invoice_no": num, "lines": [{"code": i.code, "qty": i.qty, "price": i.unit_price, "page": i.ev.page,
                                                          "file": bundle_of[r.sto]} for i in r.items]} for r in rs],
            {"number": num, "copies": copies_of[num]}, code_map=code_map)
        attribution.update(out["receipts"])
        for rv in out["review"]:
            unresolved.append(review_item(num, rv, "competition", len(bundle_paths)))
        for sto, v in out["receipts"].items():
            if not v["clear"]:
                unresolved.append({"reason": "بنود السند لم تُنسب للفاتورة بوضوح", "sto": sto, "invoice": num, "why": v["why"],
                                   "file": bundle_of[sto], "page": receipts[sto].pages[0] if receipts[sto].pages else None})
    review_stos = {}
    from report_writer import _ocr_review
    for o in outs:
        res_by = {r.sto: r for r in o["results"]}
        review_stos.update(_ocr_review(o["bundle"], res_by, unresolved))
    from job_pipeline import item_key
    for sto, f2, r2 in dups:                                # سند استلام بالرقم نفسه في أكثر من ربطة: ازدواج سداد محتمل
        r1, f1 = receipts[sto], bundle_of[sto]
        sig = lambda r: sorted((i.code, str(i.qty), str(i.unit_price)) for i in r.items)
        same = sig(r1) == sig(r2)
        where = "داخل الربطة" if f1 == f2 else "في أكثر من ربطة"
        u = {"reason": f"سند استلام مكرر {where}: يحتمل ازدواج السداد ويلزم مراجعته", "sto": sto,
             "invoice": r1.invoice_no, "file": f1, "page": r1.pages[0] if r1.pages else None,
             "dup_file": f2, "dup_page": r2.pages[0] if r2.pages else None,
             "why": (f"السند {sto} ورد في {f1} (ص{r1.pages[0] if r1.pages else '—'}) وفي {f2} (ص{r2.pages[0] if r2.pages else '—'}): "
                     + ("المحتوى مطابق؛ إن كان نسخة مكررة فلا يُصرف مرتين، وإن كان استلاماً ثانياً فيُوضَّح." if same else
                        "المحتوى مختلف بين الظهورين؛ رقم واحد لسندين ويلزم تحديد الصحيح.")
                     + " يُحتسب مبلغه مرة واحدة في التقرير.")}
        unresolved.append(u)
        review_stos.setdefault(sto, item_key(u))
    if not contract.items:
        unresolved.append({"reason": "بنود العقد (الملحق ب) لم تُقرأ؛ جدول الأصناف في ورقة التقرير الأخيرة غير محسوم"})
    return {"comp_no": comp_no, "rules": rules, "contract": contract, "receipts": receipts, "results": results, "order": order,
            "bundle_of": bundle_of, "pos": pos, "po_first_bundle": po_first_bundle, "inv_by_bundle": inv_by_bundle,
            "attribution": attribution, "unresolved": unresolved, "review_stos": review_stos, "bundles": [o["bundle"] for o in outs],
            "invoice_numbers": set(copies_of) | set(cit)}


# ------------------------------------------------------------------ مساعدات الكتابة
def _snapshot_styles(ws, rows, ncols):
    return {r: ([copy.copy(ws.cell(r, c)._style) for c in range(1, ncols + 1)], ws.row_dimensions[r].height) for r in rows}


def _apply_style(ws, snap, r, ncols):
    styles, h = snap
    for c in range(1, ncols + 1):
        ws.cell(r, c)._style = copy.copy(styles[c - 1])
    if h:
        ws.row_dimensions[r].height = h


def _clear(ws, from_row, ncols):
    for r in range(from_row, ws.max_row + 1):
        for c in range(1, ncols + 1):
            ws.cell(r, c).value = None


def _first_run_items(po):
    runs = _runs(po.pages)
    first = set(runs[0]) if runs else set()
    return [i for i in po.items if i.ev.page in first]


def _batch_no(fname):
    m = re.search(r"-(\d{2})-\d{3}\.pdf$", fname)
    return int(m.group(1)) if m else None                 # اسم بلا صيغة الدفعة: تبقى الخلية فارغة (لا يُختلق رقم)


def _bundle_no(fname):
    m = re.search(r"-(\d{3})\.pdf$", fname)
    return int(m.group(1)) if m else None


def write_competition_report(bundle_paths, contract_path, out_path, code_map=None, workers=1, progress=None, data=None, resolved_keys=None, exclude_stos=None):
    data = data or collect_competition(bundle_paths, contract_path, code_map, workers, progress)
    from job_pipeline import item_key
    open_items = [u for u in data['unresolved'] if item_key(u) not in (resolved_keys or set())]
    rules, contract = data["rules"], data["contract"]
    comp = data["comp_no"]
    order = [s for s in data["order"] if s not in set(exclude_stos or ())]      # السند الذي أكّد المراجع أن فيه مشكلة لا يدخل التقرير
    data = {**data, "order": order}
    wb = openpyxl.load_workbook(T.writer_template_path("competition"))
    ws1, ws2, ws3, ws4 = (wb[n] for n in ("تقرير السندات", "تفصيل الأصناف", "مطابقة التعاميد", "تقرير المنافسة والأصناف"))
    n1, n2, n3 = ws1.max_column, ws2.max_column, ws3.max_column
    sn1, sn2, sn3 = (_snapshot_styles(ws, [6], ws.max_column)[6] for ws in (ws1, ws2, ws3))
    for ws in (ws1, ws2, ws3):
        _clear(ws, 6, ws.max_column)

    batches = sorted({_batch_no(f) for f in data["bundle_of"].values()} - {None})
    ws1["A3"] = (f"المنافسة {comp} | {(str(len(batches)) + " دفعات") if batches else "دفعات غير مرقمة"} | أرقام الأصناف والكميات والأسعار متقابلة بالترتيب نفسه؛ "
                 "تفاصيل كل صنف في الشيت التالي")
    # ---------------- ورقة السندات + ورقة الأصناف
    r1 = r2 = 6
    contract_items = {i.code: i for i in contract.items}
    inv_numbers = set()
    for sto in order:
        rec, res = data["receipts"][sto], data["results"][sto]
        fname = data["bundle_of"][sto]
        att = data["attribution"].get(sto, {})
        pairs = att.get("pairs") or []
        inv_b = data["inv_by_bundle"][fname].get(rec.invoice_no)
        if rec.invoice_no:
            inv_numbers.add(rec.invoice_no)
        groups = [sorted(v) for v in inv_b.copy_pages.values()] if inv_b else []
        grp = _pick_after(groups, min(rec.pages)) if inv_b and rec.pages else []
        copy_id = next((k for k, v in inv_b.copy_pages.items() if sorted(v) == grp), 0) if inv_b else 0
        po = data["pos"].get(rec.pon)
        po_b = next((p for p in next(b for b in data["bundles"] if b.file == fname).pos if p.number == rec.pon), None)
        po_runs = _runs(po_b.pages) if po_b else []
        po_grp = po_runs[0] if po_runs else []
        taxable = [it["taxable"] for it in res.items]
        held = sto in data["review_stos"] and data["review_stos"][sto] not in (resolved_keys or set())
        due = rec.due or max((i.due for i in (po.items if po else []) if i.due), default=None)
        row_vals = [comp, _batch_no(fname), _bundle_no(fname), sto, _date(rec.date), rec.pon, rec.invoice_no,
                    _date(inv_b.date) if inv_b and inv_b.date else None,
                    "\n".join(it["code"] for it in res.items), "\n".join(str(_fmt_qty(it["qty"])) for it in res.items),
                    "\n".join(f"{it['price']:.2f}" for it in res.items), _f(res.net), _f(res.commission),
                    "نعم" if res.vat > 0 else "لا",
                    None if res.vat == 0 and not any(taxable) else ("كامل السند" if all(taxable) else "جزئي"),
                    _f(res.vat), "نعم" if res.penalty > 0 else "لا", _f(res.penalty), "لا", 0,
                    _f(res.net + res.commission), _f(res.total), _f(res.total - res.penalty), _f(res.payable),
                    None if held else ("سليم محاسبياً" if res.ok else "غير سليم محاسبياً"),
                    None if held else ("؛ ".join(i.text for i in res.issues) if res.issues else None),
                    fname, SEP.join(str(p) for p in sorted(rec.pages)), SEP.join(str(p) for p in sorted(inv_b.pages)) if inv_b else None,
                    SEP.join(str(p) for p in po_grp) if po_grp else None, _date(due), res.late_days]
        _apply_style(ws1, (sn1[0], sn1[1]), r1, n1)
        for c, v in enumerate(row_vals, 1):
            ws1.cell(r1, c, v)
        r1 += 1
        # ---- الأصناف
        for k, it in enumerate(res.items):
            ritem = rec.items[k]
            pair = next((p for p in pairs if p[0]["code"] == it["code"] and p[0]["qty"] == it["qty"]), None)
            inv_row = pair[1] if pair else None
            inv_page = inv_row.get("page") if inv_row else None
            expiry = inv_row.get("expiry") if inv_row else None
            if inv_row and inv_b:
                same = _same_row(inv_b, inv_row, copy_id)
                if same is not None:
                    inv_page, expiry = same.ev.page, same.expiry or expiry
            ci = contract_items.get(it["code"])
            line_ok = (inv_row is not None and ritem.line_total is not None and abs(ritem.line_total - inv_row["amount"]) <= D("0.005")
                       and ci is not None and ci.price == ritem.unit_price)
            vals = [comp, _batch_no(fname), sto, rec.pon, rec.invoice_no, it["code"], ci.description if ci else None, ritem.unit,
                    _fmt_qty(it["qty"]), _f(ritem.unit_price), _f(ci.price) if ci else None, _f(it["net"]), _f(it["comm"]),
                    float(rules.vat_rate) if it["vat"] > 0 else 0, _f(it["vat"]), _f(it["total"]),
                    _f(inv_row["net"]) if inv_row else None, _f(inv_row["nupco"]) if inv_row else None,
                    _f(inv_row["vat"]) if inv_row else None, _f(inv_row["amount"]) if inv_row else None,
                    _date(expiry), ritem.ev.page, SEP.join(str(p) for p in sorted(inv_b.pages)) if inv_b else None, (("مطابق" if line_ok else "غير مطابق") if inv_row is not None else None), fname]
            _apply_style(ws2, (sn2[0], sn2[1]), r2, n2)
            for c, v in enumerate(vals, 1):
                ws2.cell(r2, c, v)
            r2 += 1
    last1, last2 = r1 - 1, r2 - 1
    for ws, last in ((ws1, last1), (ws2, last2)):
        for r in range(last + 1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                ws.cell(r, c).value = None
        if ws.max_row > last:
            ws.delete_rows(last + 1, ws.max_row - last)

    # ---------------- مطابقة التعاميد (تعميد مستقل واحد، التكرار داخل الربطات لا يُحتسب)
    by_po = defaultdict(list)
    for sto in order:
        by_po[data["receipts"][sto].pon].append(sto)
    r3 = 6
    for pon in sorted(data["pos"]):
        po = data["pos"][pon]
        po_items = _first_run_items(po)
        stos = by_po.get(pon, [])
        po_units = sum((i.qty_ordered for i in po_items), D(0))
        po_net = sum((i.total for i in po_items), D(0))
        po_comm = engine.r2(po_net * rules.commission)
        po_vat = sum((i.vat for i in po_items), D(0))
        po_tot = sum((i.final for i in po_items), D(0))
        rc_units = sum((i.qty for s in stos for i in data["receipts"][s].items), D(0))
        rc_net = sum((data["results"][s].net for s in stos), D(0))
        rc_comm = sum((data["results"][s].commission for s in stos), D(0))
        rc_vat = sum((data["results"][s].vat for s in stos), D(0))
        rc_tot = sum((data["results"][s].total for s in stos), D(0))
        ok = (po_units == rc_units and po_net == rc_net and abs(po_comm - rc_comm) <= D("0.005")
              and abs(po_vat - rc_vat) <= D("0.005") and abs(po_tot - rc_tot) <= D("0.005"))
        due = max((i.due for i in po_items if i.due), default=None)
        vals = [comp, _batch_no(data["po_first_bundle"][pon]), pon, _date(po.release_date), _date(due), len(po_items), len(stos),
                _fmt_qty(po_units), _fmt_qty(rc_units), _f(po_net), _f(rc_net), _f(po_comm), _f(rc_comm), _f(po_vat), _f(rc_vat),
                _f(po_tot), _f(rc_tot), "مطابق" if ok else "غير مطابق"]
        _apply_style(ws3, (sn3[0], sn3[1]), r3, n3)
        for c, v in enumerate(vals, 1):
            ws3.cell(r3, c, v)
        r3 += 1
    for r in range(r3, ws3.max_row + 1):
        for c in range(1, ws3.max_column + 1):
            ws3.cell(r, c).value = None
    if ws3.max_row >= r3:
        ws3.delete_rows(r3, ws3.max_row - r3 + 1)

    _write_summary_sheet(ws4, data, batches, last1, last2, len(inv_numbers), bool(open_items))
    n_f, f_err = RS.finalize(wb, out_path)                 # تنسيق المراجعة + حفظ + نتائج الصيغ المخزّنة
    T.assert_matches_template(out_path, "competition")
    return out_path, data["unresolved"]


def _write_summary_sheet(ws, data, batches, last1, last2, n_invoices, has_open):
    comp, rules, contract = data["comp_no"], data["rules"], data["contract"]
    ncols = ws.max_column
    snap = _snapshot_styles(ws, range(1, 44), ncols)
    for r in range(6, 44):
        for c in range(1, ncols + 1):
            ws.cell(r, c).value = None
    ni, nb = len(contract.items), len(batches)
    it0 = 20
    it_end = it0 + ni - 1
    tot = it0 + ni
    bh = tot + 3
    b0 = bh + 1
    b_end = b0 + nb - 1
    btot = b_end + 1
    close = btot + 2
    note = close + 1
    # أنماط الصفوف: ننسخ من صف القالب النموذجي لكل نوع (نفس ترتيب القالب)
    S1, S2 = "'تقرير السندات'!", "'تفصيل الأصناف'!"

    def style_rows(frm, to, tpl_row):
        for r in range(frm, to + 1):
            _apply_style(ws, snap[tpl_row], r, ncols)

    for r in range(6, 18):
        _apply_style(ws, snap[r], r, ncols)
    _apply_style(ws, snap[18], 18, ncols)
    _apply_style(ws, snap[19], 19, ncols)
    style_rows(it0, it_end, 20)
    _apply_style(ws, snap[30], tot, ncols)
    for r in range(tot + 1, bh):
        _apply_style(ws, snap[31], r, ncols)
    _apply_style(ws, snap[33], bh, ncols)
    style_rows(b0, b_end, 34)
    _apply_style(ws, snap[40], btot, ncols)
    _apply_style(ws, snap[41], btot + 1, ncols)
    _apply_style(ws, snap[42], close, ncols)
    _apply_style(ws, snap[43], note, ncols)
    if ws.max_row > note:
        ws.delete_rows(note + 1, ws.max_row - note)

    ws["A3"] = f"المنافسة {comp} | {contract.supplier or ''} | {(str(len(batches)) + " دفعات") if batches else "دفعات غير مرقمة"} | المصدر: عقد المنافسة.pdf والربطات والتعاميد"
    ct = contract.totals
    lines = [(6, "صافي مبلغ المنافسة", "L", ct.get("net")), (7, "إجمالي عمولة نوبكو", "M", ct.get("nupco")),
             (8, "إجمالي ضريبة القيمة المضافة", "P", ct.get("vat")),
             (9, "الإجمالي بعد عمولة نوبكو والضريبة", "V", ct.get("total"))]
    for r, label, col, cval in lines:
        ws.cell(r, 1, label)
        ws.cell(r, 2, f"=ROUND(SUM({S1}{col}6:{col}{last1}),2)")
        ws.cell(r, 3, _f(cval))
        if cval is not None:
            ws.cell(r, 4, f"=ROUND(B{r}-C{r},2)")
            ws.cell(r, 5, f'=IF(D{r}=0,"مطابق","غير مطابق")')
    ws.cell(10, 1, "إجمالي غرامات التأخير")
    ws.cell(10, 2, f"=ROUND(SUM({S1}R6:R{last1}),2)")
    ws.cell(11, 1, "إجمالي الغرامات الأخرى")
    ws.cell(11, 2, f"=ROUND(SUM({S1}T6:T{last1}),2)")
    pen_total = sum((r.penalty for r in data["results"].values()), D(0))
    if pen_total == 0:
        ws.cell(10, 5, "لا توجد غرامة موثقة")
        ws.cell(11, 5, "لا توجد غرامة موثقة")
    ws.cell(12, 1, "إجمالي المبلغ المستحق")
    ws.cell(12, 2, f"=ROUND(SUM({S1}X6:X{last1}),2)")
    ws.cell(12, 3, _f(ct.get("total")))
    if ct.get("total") is not None:
        ws.cell(12, 4, "=ROUND(B12-C12,2)")
        ws.cell(12, 5, '=IF(D12=0,"مطابق","غير مطابق")')
    ws.cell(14, 1, "عدد السندات")
    ws.cell(14, 2, len(data["order"]))
    ws.cell(14, 3, "عدد الفواتير")
    ws.cell(14, 4, n_invoices)
    ws.cell(15, 1, "عدد التعاميد")
    ws.cell(15, 2, len(data["pos"]))
    ws.cell(15, 3, "الربطات المفحوصة")
    ws.cell(15, 4, len(data["bundles"]))
    ws.cell(17, 1, "الأصناف المستلمة - العدد الإجمالي")
    ws.cell(17, 2, f"=SUM(D{it0}:D{it_end})" if ni else 0)
    ws.cell(17, 3, "الأصناف غير المستلمة - العدد الإجمالي")
    ws.cell(17, 4, f"=SUM(E{it0}:E{it_end})" if ni else 0)
    heads = ["رقم الصنف", "وصف الصنف", "الكمية في العقد", "الكمية المستلمة", "الكمية غير المستلمة", "كمية زائدة", "سعر الوحدة",
             "صافي المستلم", "عمولة نوبكو", "ضريبة القيمة المضافة", "إجمالي المستلم", "نسبة الضريبة", "حالة الصنف"]
    for c, h in enumerate(heads, 1):
        ws.cell(19, c, h)
    rng = lambda col: f"{S2}${col}$6:${col}${last2}"
    for k, it in enumerate(contract.items):
        r = it0 + k
        ws.cell(r, 1, it.code)
        ws.cell(r, 2, it.description)
        ws.cell(r, 3, _fmt_qty(it.qty))
        ws.cell(r, 4, f"=SUMIF({rng('F')},A{r},{rng('I')})")
        ws.cell(r, 5, f"=MAX(C{r}-D{r},0)")
        ws.cell(r, 6, f"=MAX(D{r}-C{r},0)")
        ws.cell(r, 7, _f(it.price))
        ws.cell(r, 8, f"=ROUND(SUMIF({rng('F')},A{r},{rng('L')}),2)")
        ws.cell(r, 9, f"=ROUND(SUMIF({rng('F')},A{r},{rng('M')}),2)")
        ws.cell(r, 10, f"=ROUND(SUMIF({rng('F')},A{r},{rng('O')}),2)")
        ws.cell(r, 11, f"=ROUND(SUM(H{r}:J{r}),2)")
        ws.cell(r, 12, float(it.vat_rate))
        ws.cell(r, 13, f'=IF(AND(C{r}=D{r},E{r}=0,F{r}=0),"مكتمل","غير مكتمل")')
    if ni:
        ws.cell(tot, 1, "الإجمالي")
        for col in "CDEF":
            ws[f"{col}{tot}"] = f"=SUM({col}{it0}:{col}{it_end})"
        for col in "HIJK":
            ws[f"{col}{tot}"] = f"=ROUND(SUM({col}{it0}:{col}{it_end}),2)"
    for c, h in enumerate(["رقم الدفعة", "عدد السندات", "الصافي", "عمولة نوبكو", "الضريبة", "الإجمالي", "إجمالي الدفعة في العقد",
                           "الفرق", "نتيجة المطابقة"], 1):
        ws.cell(bh, c, h)
    B1 = "'تقرير السندات'!$B$6:$B$" + str(last1)
    col1 = lambda col: f"{S1}${col}$6:${col}${last1}"
    for k, b in enumerate(batches):
        r = b0 + k
        ws.cell(r, 1, b)
        ws.cell(r, 2, f"=COUNTIF({B1},A{r})")
        ws.cell(r, 3, f"=ROUND(SUMIF({B1},A{r},{col1('L')}),2)")
        ws.cell(r, 4, f"=ROUND(SUMIF({B1},A{r},{col1('M')}),2)")
        ws.cell(r, 5, f"=ROUND(SUMIF({B1},A{r},{col1('P')}),2)")
        ws.cell(r, 6, f"=ROUND(SUMIF({B1},A{r},{col1('V')}),2)")
        bc = contract.batches.get(b)
        if bc is not None:
            ws.cell(r, 7, _f(bc["total"]))
            ws.cell(r, 8, f"=ROUND(F{r}-G{r},2)")
            ws.cell(r, 9, f'=IF(H{r}=0,"مطابق","غير مطابق")')
    ws.cell(btot, 1, "إجمالي الدفعات")
    ws.cell(btot, 2, f"=SUM(B{b0}:B{b_end})")
    for col in "CDEFGH":
        ws[f"{col}{btot}"] = f"=ROUND(SUM({col}{b0}:{col}{b_end}),2)"
    ws.cell(close, 1, "حالة إقفال المنافسة")
    formula = (f'=IF(AND(D6=0,D7=0,D8=0,D9=0,D12=0,E{tot}=0,F{tot}=0,H{btot}=0,'
               f'COUNTIF({S1}Y6:Y{last1},"غير سليم محاسبياً")=0),"مقفل حسابياً","يتطلب مراجعة")')
    # غير المحسوم لا يُغلق: القيمة الحرفية «يتطلب مراجعة» (من القالب) بدل الصيغة التي لا ترى الخلايا الفارغة
    needs = has_open or not contract.totals or not contract.batches or not ni
    ws.cell(close, 2, "يتطلب مراجعة" if needs else formula)
    ws.cell(note, 1, "ملاحظة")
    ws.cell(note, 2, "بيانات تدريبية غير موقعة؛ الحكم على اتساق الملفات المقدمة.")
