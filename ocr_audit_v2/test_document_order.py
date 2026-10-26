"""Paper order regression, including generated workbook rows in all receipt sheets."""
from pathlib import Path
from types import SimpleNamespace as N
from unittest.mock import patch
import json,os,tempfile
import openpyxl
from document_order import ordered_receipts,ordered_bundle_paths,batch_receipt_ids,competition_receipt_ids

def receipt(sto,pages,top=None):
    items=[] if top is None else [N(ev=N(page=min(pages),box=(0,top,100,top+20)))]
    return N(sto=sto,pages=pages,items=items)

def units():
    rows=[receipt('STO0000000001',[85]),receipt('STO0000000900',[3,4]),receipt('STO0000000100',[12])]
    expected=['STO0000000900','STO0000000100','STO0000000001']
    assert [r.sto for r in ordered_receipts(rows)]==expected
    same=[receipt('Z',[1]),receipt('A',[1])]
    assert [r.sto for r in ordered_receipts(same)]==['Z','A']
    same=[receipt('Z',[1],100),receipt('A',[1],30)]
    assert [r.sto for r in ordered_receipts(same)]==['A','Z']
    data={'receipts':{r.sto:r for r in rows},'results':{r.sto:None for r in rows}}
    assert batch_receipt_ids(data)==expected
    assert batch_receipt_ids(data,[expected[1]])==[expected[0],expected[2]]
    assert ordered_bundle_paths(['PAY-730-10.pdf','PAY-730-2.pdf','PAY-730-1.pdf'])==['PAY-730-1.pdf','PAY-730-2.pdf','PAY-730-10.pdf']
    data['bundle_of']={rows[0].sto:'PAY-730-2.pdf',rows[1].sto:'PAY-730-2.pdf',rows[2].sto:'PAY-730-10.pdf'}
    assert competition_receipt_ids(data)==[rows[1].sto,rows[0].sto,rows[2].sto]
    print('7 paper sequence assertions passed')

def workbook_order(path,expected):
    wb=openpyxl.load_workbook(path,data_only=False);checked=[]
    for ws in wb.worksheets:
        heading=next(((cell.row,cell.column) for row in ws.iter_rows(min_row=1,max_row=8)
                      for cell in row if cell.value=='رقم السند'),None)
        if heading is None:continue
        row,col=heading;seen=[]
        for r in range(row+1,ws.max_row+1):
            value=ws.cell(r,col).value
            if isinstance(value,str) and value.startswith('STO') and value not in seen:seen.append(value)
        assert seen==expected,(ws.title,seen,expected)
        checked.append(ws.title)
    assert len(checked)>=2,checked
    return checked

def real_workbooks():
    import report_template as T,report_writer as W,competition_writer as C
    root=Path(T.ROOT);folder=next((root/'المنافسات').glob('*MED-2026-318'))
    contract=next((folder/'العقد').glob('*.pdf'))
    files=ordered_bundle_paths([str(p) for p in folder.glob('MED-2026-318-*.pdf')])[:2]
    assert len(files)==2
    out=Path(os.environ.get('ORDER_TEST_OUTPUT') or tempfile.mkdtemp(prefix='paper_order_'));out.mkdir(parents=True,exist_ok=True)
    batch=W.collect_batch(files[0],str(contract))
    # Use the source bundle sequence as an independent expectation, not the helper.
    expected=list(dict.fromkeys(r.sto for r in batch['bundle'].receipts))
    # Writer must recover paper order even if a cache returns reversed dict order.
    batch['receipts']=dict(reversed(list(batch['receipts'].items())))
    batch['results']=dict(reversed(list(batch['results'].items())))
    with patch.object(W,'collect_batch',return_value=batch):W.write_batch_report(files[0],str(contract),str(out/'batch_order.xlsx'))
    batch_sheets=workbook_order(out/'batch_order.xlsx',expected)
    assert not T.validate(str(out/'batch_order.xlsx'),'batch')
    competition=C.collect_competition(list(reversed(files)),str(contract),workers=1)
    expected=[]
    for bundle in competition['bundles']:
        for r in bundle.receipts:
            if r.sto not in expected:expected.append(r.sto)
    competition['order'].reverse()
    competition['receipts']=dict(reversed(list(competition['receipts'].items())))
    C.write_competition_report(files,str(contract),str(out/'competition_order.xlsx'),data=competition)
    competition_sheets=workbook_order(out/'competition_order.xlsx',expected)
    assert not T.validate(str(out/'competition_order.xlsx'),'competition')
    result={'batch_receipts':len(batch['receipts']),'competition_receipts':len(expected),
            'batch_sheets':batch_sheets,'competition_sheets':competition_sheets,'template_checks':'passed',
            'source_files':files,'scope':'Order and template regression only, not accuracy or closure evaluation'}
    (out/'result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps(result,ensure_ascii=False))

if __name__=='__main__':
    units();real_workbooks()
